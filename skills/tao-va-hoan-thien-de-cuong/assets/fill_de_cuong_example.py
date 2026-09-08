# -*- coding: utf-8 -*-
import copy
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from docx.shared import RGBColor

TEMPLATE="assets/Mau_de_cuong_2025.docx"   # mẫu chuẩn 2025 đóng gói kèm skill
OUTDIR="/mnt/c/ShareWork/02_Claude_Projects/TS_QLKT/QLY801_TaiLieu"

# ---------- helpers ----------
def clear_runs(p):
    for r in list(p.runs): r._element.getparent().remove(r._element)
def set_text(p, text, bold=False):
    clear_runs(p); r=p.add_run(text); r.bold=bold; return r
def set_label_value(p, label, value):
    clear_runs(p); p.add_run(label+': ').bold=True; p.add_run(value)
def insert_par_after(p, text, bold=False):
    el=copy.deepcopy(p._element); p._element.addnext(el)
    np=Paragraph(el,p._parent); set_text(np,text,bold); return np
def remove_par(p): p._element.getparent().remove(p._element)
def set_cell(cell, text, bold=False):
    cell.text=''; set_text(cell.paragraphs[0], text, bold)
def set_cell_lines(cell, lines):
    cell.text=''
    set_text(cell.paragraphs[0], lines[0], bold=True)  # dòng đầu = tên chương, in đậm
    for ln in lines[1:]:
        p=cell.add_paragraph(); set_text(p, ln)
def trim_table(tbl, keep):
    while len(tbl.rows)>keep: tbl._tbl.remove(tbl.rows[-1]._tr)

def is_red(run):
    try:
        rgb=run.font.color.rgb
    except Exception:
        rgb=None
    if rgb is None: return False
    R,G,B=rgb[0],rgb[1],rgb[2]
    return R>=0x90 and G<=0x55 and B<=0x55

def strip_red_and_footnotes(doc):
    # xóa run màu đỏ (chú thích) + footnote reference; NHƯNG giữ dòng "Kèm theo QĐ"/Decision
    KEEP=['Kèm theo','QĐ-ĐHNT','Hiệu trưởng','Attached to Decision','QD-ĐHNT','President of Foreign']
    def rm_footnote_refs(p):
        for run_el in p._element.findall(qn('w:r')):
            if run_el.find(qn('w:footnoteReference')) is not None:
                run_el.getparent().remove(run_el)
    def proc_par(p):
        if any(k in p.text for k in KEEP):
            rm_footnote_refs(p); return   # giữ nguyên dòng QĐ (kể cả chữ đỏ)
        for r in list(p.runs):
            if is_red(r): r._element.getparent().remove(r._element)
        rm_footnote_refs(p)
    for p in doc.paragraphs: proc_par(p)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs: proc_par(p)

def clean_paren_heading(doc, prefixes):
    # với các heading hướng dẫn: bỏ phần "(...)" chỉ dẫn, giữ phần chính
    import re
    for p in doc.paragraphs:
        t=p.text.strip()
        for pre in prefixes:
            if t.startswith(pre):
                base=re.sub(r'\s*\(.*$','',t)  # bỏ từ dấu ( tới hết
                if base!=t: set_text(p, base, bold=p.runs[0].bold if p.runs else False)
                break

# ---------- content ----------
TITLE_VN="Thiết kế nghiên cứu trong kinh tế và quản lý (Research Design in Economics and Management)"
TITLE_EN="Research Design in Economics and Management (Thiết kế nghiên cứu trong kinh tế và quản lý)"
CODE="QLY801"
DEPT_VN="Bộ môn Kinh tế và Quản lý"; DEPT_EN="Department of Economics and Governance"
FAC_VN="Viện Kinh tế và Kinh doanh quốc tế"; FAC_EN="School of Economics and International Business"
CRED_VN="2 tín chỉ"; CRED_EN="2 credits"
PRE_VN="Phương pháp luận nghiên cứu (PPH804)"; PRE_EN="Research Methodology (PPH804)"
DESC_VN=["Học phần cung cấp nền tảng lý thuyết và thực tiễn về thiết kế nghiên cứu, tập trung vào ba phương pháp chính: định tính, định lượng và phương pháp hỗn hợp. Học phần hướng dẫn nghiên cứu sinh phát triển kỹ năng xây dựng đề cương nghiên cứu, xác định vấn đề nghiên cứu, lựa chọn phương pháp phù hợp, thu thập và phân tích dữ liệu, cũng như trình bày kết quả nghiên cứu một cách khoa học và hiệu quả; chú trọng ứng dụng trong lĩnh vực kinh tế và quản lý và tuân thủ các tiêu chuẩn đạo đức trong nghiên cứu.",
 "Phần nền tảng triết học và phương pháp luận của nghiên cứu đã được trang bị ở học phần Phương pháp luận nghiên cứu; ở học phần này các nội dung đó chỉ được nhắc lại ngắn gọn rồi tập trung vào khâu thiết kế và viết đề cương nghiên cứu cụ thể."]
DESC_EN=["This course provides theoretical and practical foundations of research design, focusing on three core approaches: qualitative, quantitative, and mixed methods. It guides doctoral students in building a research proposal, identifying the research problem, selecting an appropriate approach, collecting and analysing data, and presenting research results scientifically, with applications in economics and management and adherence to research ethics.",
 "The philosophical and methodological foundations are covered in the Research Methodology course and are only briefly revisited here, with the focus on designing and writing a concrete research proposal."]
CLO_VN=["Đề xuất được thiết kế nghiên cứu trong lĩnh vực quản lý kinh tế và chính sách.",
 "Thực hiện được nghiên cứu độc lập trong lĩnh vực quản lý kinh tế và chính sách.",
 "Có kỹ năng phát hiện khoảng trống nghiên cứu nhằm hình thành thiết kế nghiên cứu trong kinh tế và quản lý.",
 "Phối hợp kỹ năng phân tích, tổng hợp nhằm đưa ra thiết kế nội dung nghiên cứu trong kinh tế và quản lý.",
 "Hình thành năng lực ra quyết định khi thiết kế nghiên cứu trong lĩnh vực kinh tế và quản lý.",
 "Hình thành năng lực dẫn dắt chuyên môn, thiết kế được các nghiên cứu trong lĩnh vực quản lý kinh tế."]
CLO_EN=["Propose research designs in the field of economic management and policy.",
 "Conduct independent research in the field of economic management and policy.",
 "Develop skills to identify research gaps in order to form a research design in economics and management.",
 "Combine analytical and synthesis skills to develop a research design in economics and management.",
 "Develop decision-making skills when designing research in economics and management.",
 "Develop professional leadership skills to design research in economic management."]
REF_REQ_VN="[1] Creswell, J. W., & Creswell, J. D. (2026). Thiết kế nghiên cứu (Research Design) (Phạm Minh Vương dịch, ấn bản thứ 6). IRED Books - Nhà xuất bản Tổng hợp Thành phố Hồ Chí Minh."
REF_REQ_EN="[1] Creswell, J. W., & Creswell, J. D. (2023). Research Design: Qualitative, Quantitative, and Mixed Methods Approaches (6th ed.). SAGE Publications."
REF_OPT_VN=["[2] Bukve, O. (2019). Designing Social Science Research. Palgrave Macmillan.",
 "[3] Blaikie, N., & Priest, J. (2019). Designing Social Research: The Logic of Anticipation (3rd ed.). Polity Press.",
 "[4] Maggetti, M., Radaelli, C., & Gilardi, F. (2013). Designing Research in the Social Sciences. SAGE Publications."]
REF_OPT_EN=REF_OPT_VN

# CLO -> PLO (khớp ma trận CTĐT TS QLKT: học phần QLY801 đóng góp PLO1,2,3,5,7)
CLO_PLO={1:[3],2:[1,2],3:[5],4:[5],5:[7],6:[7]}
NPLO=8
# Ánh xạ HỌC PHẦN -> PLO theo Bản mô tả CTĐT (dòng QLY801: PLO1,2,3,5,7 = mức 3)
COURSE_PLO={1:'3',2:'3',3:'3',5:'3',7:'3'}

# buổi: (n, lt, th, tl, tu, clo, vn_title, en_title, vn_lines, en_lines, vn_act, en_act)
def B(n,lt,th,tl,tu,clo,vt,et,vl,el,va,ea): return (n,lt,th,tl,tu,clo,vt,et,vl,el,va,ea)
BUOI=[
B(1,'3','0','3','6','1, 3, 5, 6','Chương 1: Lựa chọn phương pháp nghiên cứu, chiến lược viết và cân nhắc đạo đức','Ch.1: Choosing a research approach, writing strategies and ethical considerations',
 ["Chương 1: Lựa chọn phương pháp nghiên cứu, chiến lược viết và cân nhắc đạo đức","1.1 Ba cách tiếp cận nghiên cứu","1.2 Ba thành phần của một cách tiếp cận:","- Thế giới quan triết học","- Thiết kế nghiên cứu","- Phương pháp nghiên cứu","1.3 Kết nối thế giới quan, thiết kế và phương pháp","1.4 Tiêu chí lựa chọn phương pháp","1.5 Viết đề xuất nghiên cứu","1.6 Chiến lược viết","1.7 Các vấn đề đạo đức trong nghiên cứu","Tự học: đọc tài liệu [1], chương 1 và 4."],
 ["Ch.1: Choosing a research approach, writing strategies and ethical considerations","1.1 Three approaches to research","1.2 Three components of an approach:","- Philosophical worldviews","- Research designs","- Research methods","1.3 Interconnecting worldviews, designs and methods","1.4 Criteria for selecting an approach","1.5 Writing the proposal","1.6 Writing strategies","1.7 Ethical issues in research","Self-study: reference [1], Chapters 1 and 4."],
 'Thuyết giảng, thảo luận','Lecture, discussion'),
B(2,'3','0','3','6','1, 3, 4, 5','Chương 2: Tổng quan tài liệu và sử dụng lý thuyết','Ch.2: Literature review and the use of theory',
 ["Chương 2: Tổng quan tài liệu và sử dụng lý thuyết","2.1 Chọn và định đề tài nghiên cứu","2.2 Mục đích và cách tổ chức tổng quan tài liệu (định tính / định lượng / hỗn hợp)","2.3 Sử dụng tài liệu:","- Các bước tiến hành","- Tìm kiếm cơ sở dữ liệu","- Ưu tiên lựa chọn tài liệu","- Tóm tắt tài liệu","- Bản đồ tài liệu","2.4 Chuẩn trình bày trích dẫn","2.5 Định nghĩa thuật ngữ","2.6 Sử dụng lý thuyết (định lượng / định tính / hỗn hợp)","Tự học: đọc tài liệu [1], chương 2 và 3."],
 ["Ch.2: Literature review and the use of theory","2.1 Selecting a research topic","2.2 Purpose and organisation of a literature review (qualitative / quantitative / mixed)","2.3 Using the literature:","- Steps in conducting a review","- Searching databases","- Priority of the literature","- Abstracting the literature","- A literature map","2.4 Style manual use","2.5 Definition of terms","2.6 The use of theory (quantitative / qualitative / mixed)","Self-study: reference [1], Chapters 2 and 3."],
 'Thuyết giảng, thảo luận','Lecture, discussion'),
B(3,'3','0','3','6','1, 3, 4, 5','Chương 3: Giới thiệu','Ch.3: The introduction',
 ["Chương 3: Giới thiệu","3.1 Tầm quan trọng của phần giới thiệu","3.2 Bản tóm tắt cho một nghiên cứu","3.3 Giới thiệu theo ba phương pháp","3.4 Mô hình giới thiệu:","- Vấn đề nghiên cứu","- Bằng chứng từ tài liệu","- Biện minh cho vấn đề","- Khoảng trống trong tài liệu","Tự học: đọc tài liệu [1], chương 5."],
 ["Ch.3: The introduction","3.1 The importance of introductions","3.2 An abstract for a study","3.3 Qualitative, quantitative and mixed methods introductions","3.4 A model for an introduction:","- The research problem","- Evidence from the literature","- Justifying the problem","- Deficiencies in the literature","Self-study: reference [1], Chapter 5."],
 'Thuyết giảng, thảo luận','Lecture, discussion'),
B(4,'2','1','2.5','7.5','1, 3, 4, 5','Chương 4: Tuyên bố mục đích','Ch.4: The purpose statement',
 ["Chương 4: Tuyên bố mục đích","4.1 Ý nghĩa của tuyên bố mục đích","4.2 Tuyên bố mục đích định tính","4.3 Tuyên bố mục đích định lượng","4.4 Tuyên bố mục đích hỗn hợp","Thực hành: viết tuyên bố mục đích","Tự học: đọc tài liệu [1], chương 6."],
 ["Ch.4: The purpose statement","4.1 Significance of a purpose statement","4.2 A qualitative purpose statement","4.3 A quantitative purpose statement","4.4 A mixed methods purpose statement","Practice: writing a purpose statement","Self-study: reference [1], Chapter 6."],
 'Thuyết giảng, thực hành','Lecture, practice'),
B(5,'2','1','2.5','7.5','1, 2, 3, 4','Chương 5: Câu hỏi nghiên cứu và giả thuyết','Ch.5: Research questions and hypotheses',
 ["Chương 5: Câu hỏi nghiên cứu và giả thuyết","5.1 Câu hỏi nghiên cứu định tính","5.2 Câu hỏi nghiên cứu và giả thuyết định lượng","5.3 Câu hỏi và giả thuyết hỗn hợp","Thực hành: đặt câu hỏi và giả thuyết","Tự học: đọc tài liệu [1], chương 7."],
 ["Ch.5: Research questions and hypotheses","5.1 Qualitative research questions","5.2 Quantitative research questions and hypotheses","5.3 Mixed methods research questions and hypotheses","Practice: formulating questions and hypotheses","Self-study: reference [1], Chapter 7."],
 'Thuyết giảng, thực hành','Lecture, practice'),
B(6,'2','1','2.5','7.5','1, 2, 3, 4','Chương 6: Phương pháp định tính','Ch.6: Qualitative methods',
 ["Chương 6: Phương pháp định tính","6.1 Đặc điểm của nghiên cứu định tính","6.2 Các thiết kế định tính","6.3 Vai trò và tính phản tư của nhà nghiên cứu","6.4 Quy trình thu thập dữ liệu","6.5 Quy trình ghi chép dữ liệu","6.6 Quy trình phân tích dữ liệu","6.7 Diễn giải","6.8 Tính hiệu lực và độ tin cậy","6.9 Viết báo cáo nghiên cứu định tính","Tự học: đọc tài liệu [1], chương 9."],
 ["Ch.6: Qualitative methods","6.1 Characteristics of qualitative research","6.2 Qualitative designs","6.3 The researcher's role and reflexivity","6.4 Data collection procedures","6.5 Data recording procedures","6.6 Data analysis procedures","6.7 Interpretation","6.8 Validity and reliability","6.9 Writing the qualitative report","Self-study: reference [1], Chapter 9."],
 'Thuyết giảng, thực hành','Lecture, practice'),
B(7,'3','0','3','6','1, 2, 3, 4','Chương 7: Phương pháp định lượng','Ch.7: Quantitative methods',
 ["Chương 7: Phương pháp định lượng","7.1 Xác định khảo sát và thí nghiệm","7.2 Các thành phần của kế hoạch phương pháp khảo sát","7.3 Các thành phần của kế hoạch phương pháp thí nghiệm","Tự học: đọc tài liệu [1], chương 8."],
 ["Ch.7: Quantitative methods","7.1 Defining surveys and experiments","7.2 Components of a survey method plan","7.3 Components of an experimental method plan","Self-study: reference [1], Chapter 8."],
 'Thuyết giảng, thảo luận','Lecture, discussion'),
B(8,'3','0','3','7','1, 2, 3, 4','Chương 8: Quy trình phương pháp hỗn hợp','Ch.8: Mixed methods procedures',
 ["Chương 8: Quy trình phương pháp hỗn hợp","8.1 Các thành phần của quy trình hỗn hợp","8.2 Các loại thiết kế hỗn hợp:","- Hội tụ","- Giải thích tuần tự","- Khám phá tuần tự","- Thiết kế phức hợp","8.3 Quy trình nhúng thiết kế và vẽ sơ đồ","8.4 Các yếu tố quan trọng khi thiết kế (hiệu lực và đạo đức trong nghiên cứu hỗn hợp)","Trình bày và bảo vệ đề cương nghiên cứu","Tự học: đọc tài liệu [1], chương 10."],
 ["Ch.8: Mixed methods procedures","8.1 Components of mixed methods procedures","8.2 Types of mixed methods designs:","- Convergent","- Explanatory sequential","- Exploratory sequential","- Complex designs","8.3 Embedding designs and drawing diagrams","8.4 Key factors in designing (validity and ethics in mixed methods)","Proposal presentation and defense","Self-study: reference [1], Chapter 10."],
 'Thuyết giảng, thảo luận, bảo vệ đề cương','Lecture, discussion, proposal defense'),
]

def find(doc, prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix): return p
    return None

def fill_refs(doc, heading_prefix, refs):
    paras=doc.paragraphs; hi=None
    for i,p in enumerate(paras):
        if p.text.strip().startswith(heading_prefix): hi=i; break
    if hi is None: return
    bullets=[]; j=hi+1
    while j<len(paras) and (paras[j].text.strip().startswith('*') or paras[j].text.strip()==''):
        if paras[j].text.strip().startswith('*'): bullets.append(paras[j])
        j+=1
    for k,ref in enumerate(refs):
        if k<len(bullets): set_text(bullets[k],ref)
        else: insert_par_after(bullets[-1] if bullets else paras[hi], ref)
    for k in range(len(refs),len(bullets)): remove_par(bullets[k])

def add_matrix(doc, lang):
    pre = '3.2.' if lang=='vn' else '3.2.'
    p=find(doc,'3.2')
    if p is None: return
    label = 'CĐR học phần' if lang=='vn' else 'CLO'
    course_label = 'Học phần' if lang=='vn' else 'Course'
    tbl=doc.add_table(rows=1+6+1, cols=1+NPLO)   # header + 6 CLO + 1 dòng học phần
    try: tbl.style='Table Grid'
    except Exception: pass
    hdr=tbl.rows[0].cells
    set_cell(hdr[0],label,True)
    for j in range(NPLO): set_cell(hdr[1+j],'PLO%d'%(j+1),True)
    for i in range(1,7):
        cells=tbl.rows[i].cells
        set_cell(cells[0],'CLO%d'%i,True)
        for j in range(NPLO):
            set_cell(cells[1+j],'X' if (j+1) in CLO_PLO[i] else '')
    # dòng cuối: ánh xạ học phần -> PLO theo CTĐT
    crow=tbl.rows[7].cells
    set_cell(crow[0],course_label,True)
    for j in range(NPLO):
        set_cell(crow[1+j], COURSE_PLO.get(j+1,''), True)
    p._element.addnext(tbl._tbl)

def fill_common(doc, lang):
    vn=(lang=='vn')
    if vn:
        set_label_value(find(doc,'Tên học phần'),'Tên học phần',TITLE_VN)
        set_label_value(find(doc,'Mã học phần'),'Mã học phần',CODE)
        set_label_value(find(doc,'Bộ môn phụ trách'),'Bộ môn phụ trách',DEPT_VN)
        set_label_value(find(doc,'Viện/Khoa'),'Viện/Khoa',FAC_VN)
        set_label_value(find(doc,'Số tín chỉ'),'Số tín chỉ',CRED_VN)
        set_label_value(find(doc,'Điều kiện tiên quyết'),'Điều kiện tiên quyết',PRE_VN)
        set_text(find(doc,'TÊN HỌC PHẦN'),'THIẾT KẾ NGHIÊN CỨU TRONG KINH TẾ VÀ QUẢN LÝ',True)
        dp=find(doc,'(Bao gồm mục tiêu'); set_text(dp,DESC_VN[0]); insert_par_after(dp,DESC_VN[1])
        set_text(find(doc,'- CLO1'),'CLO1: '+CLO_VN[0])
        set_text(find(doc,'- CLO2'),'CLO2: '+CLO_VN[1])
        d3=find(doc,'…'); set_text(d3,'CLO3: '+CLO_VN[2]); prev=d3
        for i in range(3,6): prev=insert_par_after(prev,'CLO%d: %s'%(i+1,CLO_VN[i]))
        fill_refs(doc,'4.2. Tài liệu tham khảo bắt buộc',[REF_REQ_VN])
        fill_refs(doc,'4.3. Tài liệu tham khảo tự chọn',REF_OPT_VN)
    else:
        set_label_value(find(doc,'Course title'),'Course title',TITLE_EN)
        set_label_value(find(doc,'Course code'),'Course code',CODE)
        set_label_value(find(doc,'Department:'),'Department',DEPT_EN)
        set_label_value(find(doc,'Faculty/School'),'Faculty/School',FAC_EN)
        set_label_value(find(doc,'Credit hours'),'Credit hours',CRED_EN)
        set_label_value(find(doc,'Prerequisite'),'Prerequisite(s)',PRE_EN)
        set_text(find(doc,'COURSE TITLE'),'RESEARCH DESIGN IN ECONOMICS AND MANAGEMENT',True)
        dp=find(doc,'(Include Course Objectives'); set_text(dp,DESC_EN[0]); insert_par_after(dp,DESC_EN[1])
        set_text(find(doc,'- CLO1'),'CLO1: '+CLO_EN[0])
        set_text(find(doc,'- CLO2'),'CLO2: '+CLO_EN[1])
        d3=find(doc,'…'); set_text(d3,'CLO3: '+CLO_EN[2]); prev=d3
        for i in range(3,6): prev=insert_par_after(prev,'CLO%d: %s'%(i+1,CLO_EN[i]))
        fill_refs(doc,'4.2. Compulsory reading',[REF_REQ_EN])
        fill_refs(doc,'4.3. Optional reading',REF_OPT_EN)

def fill_tables(doc, lang, t_instr, t_51, t_52, t_assess):
    vn=(lang=='vn')
    set_cell(t_instr.rows[1].cells[0],'1')
    set_cell(t_instr.rows[1].cells[1],'TS. Nguyễn Văn A' if vn else 'Dr. Nguyen Van A')
    set_cell(t_instr.rows[1].cells[2],'giangvien@ftu.edu.vn'); set_cell(t_instr.rows[1].cells[4],'HN')
    def fmt(x): return str(int(x)) if x==int(x) else (str(x).replace('.',',') if vn else str(x))
    # 5.1 chỉ tên chương
    trim_table(t_51,2); tot=[0.0,0.0,0.0,0.0]
    for b in BUOI:
        n,lt,th,tl,tu,clo=b[0],b[1],b[2],b[3],b[4],b[5]; title=b[6] if vn else b[7]
        r=t_51.add_row().cells
        set_cell(r[0],str(n)); set_cell(r[1],title)
        set_cell(r[2],lt); set_cell(r[3],th); set_cell(r[4],tl); set_cell(r[5],tu); set_cell(r[6],clo)
        tot[0]+=float(lt); tot[1]+=float(th); tot[2]+=float(tl); tot[3]+=float(tu)
    r=t_51.add_row().cells
    set_cell(r[0],'Tổng cộng (giờ)' if vn else 'Total (hours)',True)
    set_cell(r[2],fmt(tot[0]),True); set_cell(r[3],fmt(tot[1]),True); set_cell(r[4],fmt(tot[2]),True); set_cell(r[5],fmt(tot[3]),True)
    try: r[0].merge(r[1])
    except Exception: pass
    # 5.2 mỗi buổi 5 dòng hoạt động, Nội dung xuống dòng từng ý
    trim_table(t_52,1)
    acts=(['Lý thuyết','Thực hành, thảo luận','Tiểu luận, bài tập lớn, thực tế','Tự học, chuẩn bị có hướng dẫn','Kiểm tra, đánh giá'] if vn
          else ['Lecture','Practice, discussion','Essay, assignment, fieldwork','Guided self-study','Assessment'])
    for b in BUOI:
        n,clo=b[0],b[5]; lines=b[8] if vn else b[9]; hours=[b[1],b[2],b[3],b[4],'']
        rws=[]
        for k in range(5):
            row=t_52.add_row(); rws.append(row); c=row.cells
            set_cell(c[0],str(n) if k==0 else ''); set_cell(c[1],acts[k]); set_cell(c[2],hours[k])
            if k==0: set_cell_lines(c[3],lines); set_cell(c[4],clo)
            else: set_cell(c[3],''); set_cell(c[4],'')
        for col in (0,3,4):
            m=rws[0].cells[col]
            for k in range(1,5): m=m.merge(rws[k].cells[col])
    # 7 đánh giá
    trim_table(t_assess,1)
    rows=( [('Đánh giá quá trình','Chuyên cần','Số lần có mặt và mức độ tham gia thảo luận','5, 6','10%'),
            ('Đánh giá quá trình','Đánh giá định kỳ (giữa kỳ)','Bài tập/tiểu luận về một cấu phần thiết kế nghiên cứu','1, 3, 4','30%'),
            ('Đánh giá tổng kết','Bài nghiên cứu cuối kỳ','Đề cương/thiết kế nghiên cứu hoàn chỉnh; tiêu chí: rõ vấn đề, phù hợp phương pháp, chặt chẽ lập luận, tuân thủ đạo đức','1, 2, 3, 4','60%')] if vn
          else [('Process assessment','Attendance','Class attendance and participation','5, 6','10%'),
            ('Process assessment','Periodic (mid-term)','Assignment/essay on a component of research design','1, 3, 4','30%'),
            ('Summative assessment','Final research paper','A complete research proposal/design; criteria: clarity, appropriateness, rigour, ethics','1, 2, 3, 4','60%')])
    nc=len(t_assess.columns)
    for d in rows:
        c=t_assess.add_row().cells
        set_cell(c[0],d[0]); set_cell(c[1],d[1]); set_cell(c[2],d[2])
        if nc>=6: set_cell(c[3],''); set_cell(c[4],d[3]); set_cell(c[5],d[4])
        else: set_cell(c[nc-2],d[3]); set_cell(c[nc-1],d[4])
    c=t_assess.add_row().cells
    set_cell(c[nc-2],'Tổng' if vn else 'Total',True); set_cell(c[nc-1],'100%',True)

def split_inplace(doc, lang):
    body=doc.element.body; kids=list(body.iterchildren()); syll=None
    for i,ch in enumerate(kids):
        if ch.tag==qn('w:p') and Paragraph(ch,doc).text.strip()=='SYLLABUS': syll=i; break
    if lang=='vn':
        for ch in kids[syll:]:
            if ch.tag==qn('w:sectPr'): continue
            body.remove(ch)
    else:
        for ch in kids[:syll]: body.remove(ch)

def trim_instr(doc):
    t=doc.tables[0]
    while len(t.rows)>2: t._tbl.remove(t.rows[-1]._tr)

def build(lang, outpath, add_sig=False):
    doc=Document(TEMPLATE)
    split_inplace(doc,lang)
    clean_paren_heading(doc,['4.1.','4.4.','3.2.','4.1','4.4'] if lang=='vn' else ['4.1.','4.4.','3.2.'])
    fill_common(doc,lang)
    fill_tables(doc,lang,doc.tables[0],doc.tables[1],doc.tables[2],doc.tables[3])
    add_matrix(doc,lang)
    trim_instr(doc)
    strip_red_and_footnotes(doc)
    if add_sig:
        doc.add_paragraph("")
        sig=doc.add_table(rows=1,cols=2)
        sig.rows[0].cells[0].paragraphs[0].add_run("TRƯỞNG KHOA").bold=True
        sig.rows[0].cells[1].paragraphs[0].add_run("TRƯỞNG BỘ MÔN").bold=True
        for c in sig.rows[0].cells: c.paragraphs[0].alignment=1
    doc.save(outpath)
    print(lang,"saved; tables=",len(doc.tables))

build('vn',OUTDIR+'/DeCuong_QLY801_VN_2025.docx',add_sig=True)
build('en',OUTDIR+'/Syllabus_QLY801_EN_2025.docx',add_sig=False)
print("DONE")
