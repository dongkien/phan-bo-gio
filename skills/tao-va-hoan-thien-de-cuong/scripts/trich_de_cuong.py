#!/usr/bin/env python3
"""Trích xuất một đề cương FTU (mẫu 2025, .docx) thành bản tóm tắt có cấu trúc để RÀ CHẤT LƯỢNG:
CLO (kèm mức Bloom đoán theo động từ mở đầu), ma trận 3.2, bảng 5.1/5.2 (hoạt động dạy học theo CLO),
bảng đánh giá (hình thức theo CLO), học liệu theo mục, và bảng đối chiếu CLO -> hoạt động -> đánh giá.

Dùng:  python3 trich_de_cuong.py DeCuong.docx            # in Markdown
       python3 trich_de_cuong.py DeCuong.docx --json     # in JSON
Cần: pip install python-docx. Không nối mạng, không sửa file.
Script chỉ TRÍCH và ĐỐI CHIẾU cơ học; việc đánh giá "có phù hợp không" là của người/Claude đọc bản trích này
cùng rubric ở references/ra_soat_chat_luong.md."""
import sys, re, json, unicodedata
from docx import Document

NFC = lambda s: unicodedata.normalize('NFC', s or '')
low = lambda s: NFC(s).lower().strip()

# ---- mức năng lực theo động từ mở đầu (Phụ lục 1-3 Hướng dẫn CĐR FTU + Bloom tiếng Anh) ----
LEVELS = {
 'K1 Nhớ': ['nhận diện','nhận biết','nhớ lại','nhớ','nêu tên','nêu','gọi tên','gán tên','trích dẫn','liệt kê','nối','định nghĩa','nhắc lại','chỉ ra','kể tên','identify','recognize','recognise','recall','retrieve','name','label','quote','list','match','define','repeat','state'],
 'K2 Hiểu': ['phát biểu','mô tả','dịch','phân loại','tóm kết','tóm tắt','so sánh','khái quát hóa','khái quát hoá','khái quát','giải thích','minh họa','minh hoạ','phân biệt','sắp xếp lại','sắp xếp','viết lại','trình bày','diễn giải','diễn đạt','describe','translate','classify','sort','summarize','summarise','compare','generalize','generalise','explain','illustrate','distinguish','rearrange','reorder','rewrite','present','interpret','discuss','outline','paraphrase'],
 'K3 Áp dụng': ['vận dụng','áp dụng','sử dụng','triển khai','thực hiện','vận hành','đo lường','điều chỉnh','chuyển hóa','chuyển hoá','khám phá','giải quyết','tính toán','tính','thay đổi','thiết lập','điều tra','mở rộng','tái cấu trúc','thể hiện','dự báo','dự đoán','phát hiện','ước lượng','ước tính','mô phỏng','lượng hóa','lượng hoá','xử lý','khai thác','tra cứu','thu thập','kiểm định','apply','use','utilize','utilise','execute','implement','operate','measure','adapt','modify','transfer','convert','explore','solve','calculate','compute','change','establish','investigate','expand','restructure','demonstrate','predict','forecast','detect','estimate','simulate','test','conduct','carry out'],
 'K4 Phân tích': ['phân tích','tổng hợp','chia','tìm ra','tìm','phản biện','tranh luận','sơ đồ hóa','sơ đồ hoá','tối đa hóa','tối thiểu hóa','đối chiếu','mô hình hóa','mô hình hoá','suy luận','suy ra','rút ra','chứng minh','analyze','analyse','discriminate','differentiate','synthesize','synthesise','destructure','correlate','criticize','criticise','critique','debate','characterize','characterise','diagram','maximize','maximise','minimize','minimise','examine','infer','deduce','derive','model','prove','determine'],
 'K5 Đánh giá': ['lựa chọn','chọn','đánh giá','thẩm định','thẩm tra','ngoại suy','đồng tình','phản đối','thỏa hiệp','ủng hộ','xác trị','xác nhận','xác định','xếp hạng','phán quyết','tiên đoán','đưa ra','kết luận','chẩn đoán','biện luận','lập luận','luận giải','nhận định','nhận xét','bình luận','phê phán','tư vấn','khuyến nghị','assess','evaluate','appraise','extrapolate','approve','disapprove','reconcile','support','validate','confirm','verify','grade','rank','rate','judge','conclude','diagnose','justify','argue','recommend','critically evaluate','critically analyze','critically analyse','critically assess'],
 'K6 Sáng tạo': ['viết','soạn thảo','soạn','báo cáo','thiết kế','xây dựng','kết hợp','sáng tác','chuẩn hóa','chuẩn hoá','hệ thống hóa','hệ thống hoá','lập kế hoạch','lập công thức','lập trình','lập','hoạch định','phát triển','lắp ráp','cải thiện','phát minh','sáng chế','tổ chức lại','tổ chức','sửa lỗi','sáng tạo','tạo ra','tạo','kiến tạo','đề xuất','đề ra','write','report','design','build','construct','combine','incorporate','compose','standardize','standardise','hypothesize','hypothesise','develop','assemble','code','program','improve','invent','reorganize','reorganise','debug','plan','formulate','cultivate','create','produce','propose','generate','originate'],
 'S Kỹ năng': ['tuân theo','sao chép','chép lại','theo dõi','lặp lại','bắt chước','so khớp','tái hiện','nhân rộng','quan sát','thử','căn chỉnh','hiệu chỉnh','hoàn thành','tiến hành','biểu diễn','sản xuất','tái tạo','sửa chữa','sửa','phác thảo','phối hợp','giữ gìn','giữ','kiểm soát','thành thục','thành thạo','làm chủ','hoàn thiện','thay thế','đa dạng hóa','đa dạng hoá','hoàn chỉnh','chế tác','làm việc','giao tiếp','thuyết trình','quản lý','quản trị','điều hành','tích hợp','vận động','copy','duplicate','follow','imitate','reenact','replicate','observe','try','align','calibrate','complete','recreate','repair','sharpen','coordinate','balance','control','master','replace','alter','fix','vary','work','communicate','collaborate','manage','integrate','organize','organise','perform','exhibit','display','show'],
 'A Thái độ': ['tuân thủ','chú ý','lắng nghe','chấp nhận','đọc hiểu','đọc','ghi nhận','đặt câu hỏi','tham gia','thảo luận','trả lời','kể lại','thực hành','phản hồi','tôn trọng','chia sẻ','khen ngợi','ưu tiên','liên hệ','hình thành','hành động','bảo vệ','cam kết','dẫn dắt','tổng kết','phản ánh','adhere','comply','conform','attend','listen','accept','read','acknowledge','ask','participate','answer','tell','practice','practise','respond','respect','share','suggest','maintain','praise','prioritize','prioritise','relate','form','act','defend','commit','lead','engage','contribute','adopt','uphold'],
}
VAGUE = ['hiểu biết','hiểu rõ','hiểu được','hiểu','biết được','biết','nắm được','nắm vững','nắm bắt','nắm','nhận thức','có kiến thức','có hiểu biết','có khả năng','có năng lực','có kỹ năng','có thể','được trang bị','trang bị','understand','know','learn','appreciate','be familiar','be aware','comprehend','grasp','have knowledge','gain','acquire','be able to']
def starts(t, v): return t.startswith(v) and (len(t) == len(v) or not t[len(v)].isalpha())
def clo_level(text):
    t = low(re.sub(r'^[-–•*\s]*(CLO\s*\d+(\.\d+)*\s*[:.\-–]?\s*)?', '', NFC(text), flags=re.I))
    for v in VAGUE:
        if starts(t, v): return 'MƠ HỒ (' + v + ')'
    for lv, vs in LEVELS.items():
        for v in sorted(vs, key=len, reverse=True):
            if starts(t, v): return lv
    return '? (ngoài bảng)'

# ---- đọc docx ----
def grid(tb):
    g = []
    for r in tb.rows:
        row = []
        for c in r.cells:
            row.append(NFC(c.text).strip())
        g.append(row)
    return g
def dedupe(row):  # ô gộp ngang lặp text: bỏ lặp liền kề
    out = []
    for c in row:
        if not out or out[-1] != c: out.append(c)
    return out
def clo_refs(t):
    t = re.sub(r'CLO', ' ', NFC(t), flags=re.I)
    out = set()
    for a, b in re.findall(r'(\d+)\s*[-–]\s*(\d+)', t):
        out.update(range(int(a), min(int(b), 99) + 1))
    t = re.sub(r'(\d+)\s*[-–]\s*(\d+)', ' ', t)
    out.update(int(x) for x in re.findall(r'\d+', t))
    return sorted(k for k in out if k > 0)
def num(s):
    s = NFC(s).strip().replace(',', '.')
    try: return float(s)
    except: return 0.0

def extract(path):
    d = Document(path)
    paras = [NFC(p.text).strip() for p in d.paragraphs]
    out = {'file': path, 'dau_de': {}, 'mo_ta': '', 'clo': [], 'ma_tran': None, 'b51': [], 'b52': {}, 'danh_gia': [], 'hoc_lieu': {}, 'ghi_chu': []}
    # đầu đề
    for p in paras[:60]:
        m = re.match(r'^(Tên học phần|Course title|Mã học phần|Course code|Khoa phụ trách|Bộ môn phụ trách|Department|Trường/Khoa|Trường|Khoa|Viện/Khoa|Viện|Faculty/College|College|Faculty|Số tín chỉ|Credit hours|Điều kiện tiên quyết|Prerequisite\(s\))\s*:\s*(.*)$', p, flags=re.I)
        if m and m.group(2).strip(): out['dau_de'][m.group(1)] = m.group(2).strip()
    def idx(pat):
        for i, p in enumerate(paras):
            if re.match(pat, p, flags=re.I): return i
        return None
    i2, i3, i31, i32, i4, i5 = idx(r'^2\.\s'), idx(r'^3\.\s'), idx(r'^3\.1\.?\s'), idx(r'^3\.2\.?\s'), idx(r'^4\.\s'), idx(r'^5\.\s')
    if i2 is not None and i3 is not None:
        out['mo_ta'] = ' '.join(p for p in paras[i2 + 1:i3] if p and not p.startswith('('))
    if i31 is not None and i32 is not None:
        for p in paras[i31 + 1:i32]:
            if re.match(r'^-?\s*CLO\s*\d', p, flags=re.I):
                text = re.sub(r'^-?\s*CLO\s*\d+(\.\d+)*\s*[:.\-–]?\s*', '', p, flags=re.I)
                out['clo'].append({'n': len(out['clo']) + 1, 'text': text, 'muc': clo_level(text)})
    # học liệu 4.x
    if i4 is not None and i5 is not None:
        cur = None
        for p in paras[i4 + 1:i5]:
            if re.match(r'^4\.\d', p): cur = p; out['hoc_lieu'][cur] = []
            elif cur and p and not re.match(r'^\*', p): out['hoc_lieu'][cur].append(p)
    # bảng
    for tb in d.tables:
        g = grid(tb)
        if not g or not g[0]: continue
        head = ' '.join(g[0]).lower() + ' ' + (' '.join(g[1]).lower() if len(g) > 1 else '')
        if re.search(r'plo\s*\d', head) and any(re.match(r'^CLO\s*\d', r[0] if r else '', flags=re.I) or (len(r) > 1 and re.match(r'^CLO\s*\d', r[1], flags=re.I)) for r in g):
            hi = next(i for i, r in enumerate(g[:2]) if re.search(r'PLO\s*\d', ' '.join(r), flags=re.I))
            cols = [j for j, c in enumerate(g[hi]) if re.match(r'PLO', c, flags=re.I)]
            names = []
            for j in cols:  # có dòng PI dưới PLO (kiểu KTĐN) thì ghép tên PLOx/PIx.y
                nm = g[hi][j]
                if hi + 1 < len(g) and re.match(r'PI\s*\d', g[hi + 1][j], flags=re.I): nm += '/' + g[hi + 1][j]
                names.append(nm)
            rows = []
            for r in g[hi + 1:]:
                lab = next((c for c in r[:2] if re.match(r'^(CLO\s*\d|Học phần|Course|Tổng)', c, flags=re.I)), '')
                if lab and len(r) > max(cols): rows.append({'nhan': lab, 'gia_tri': {names[k]: r[j] for k, j in enumerate(cols)}})
            out['ma_tran'] = {'plo': names, 'dong': rows}
        elif re.match(r'^(Buổi|No\.?|Session)', g[0][0], flags=re.I) and ('nội dung' in head or 'content' in head) and len(g[0]) >= 6:
            for r in g:
                if re.match(r'^\d+', r[0]) and not re.match(r'^Tổng|^Total', r[0], flags=re.I):
                    nums = [c for c in r[2:] if re.match(r'^\d+([.,]\d+)?$', c)]
                    out['b51'].append({'buoi': r[0], 'noi_dung': r[1], 'hinh_thuc': r[2] if not re.match(r'^\d', r[2]) else '', 'gio': [num(x) for x in nums[:4]], 'clo': clo_refs(r[-1])})
        elif len(g[0]) == 5 and re.search(r'hoạt động|activit', head):
            cur = None
            for r in g[1:]:
                if r[0]: cur = r[0]
                if cur is None: continue
                act, gio, nd, clo = r[1], r[2], r[3], r[4]
                out['b52'].setdefault(cur, {'clo': clo_refs(clo), 'hoat_dong': []})
                if clo_refs(clo): out['b52'][cur]['clo'] = sorted(set(out['b52'][cur]['clo']) | set(clo_refs(clo)))
                out['b52'][cur]['hoat_dong'].append({'loai': act, 'gio': num(gio), 'noi_dung': nd})
        elif re.search(r'trọng số|weight|proportion', head) and re.search(r'clo|learning outcome', head):
            for r in g[1:]:
                r2 = r
                if re.match(r'^Tổng|^Total', ''.join(r2), flags=re.I): continue
                out['danh_gia'].append({'nhom': r2[0], 'hinh_thuc': r2[1] if len(r2) > 5 else '', 'noi_dung': r2[-4], 'cong_cu': r2[-3], 'clo': clo_refs(r2[-2]), 'trong_so': r2[-1]})
    # đối chiếu CLO -> buổi/hoạt động -> đánh giá
    doi_chieu = []
    for c in out['clo']:
        k = c['n']
        buoi51 = [r['buoi'] for r in out['b51'] if k in r['clo']]
        acts = {}
        gio_th = 0.0
        for b, v in out['b52'].items():
            if k in v['clo']:
                for a in v['hoat_dong']:
                    if a['gio'] > 0 or a['noi_dung']:
                        acts[a['loai']] = acts.get(a['loai'], 0) + a['gio']
                        if re.search(r'thực hành|thảo luận|practice|seminar', a['loai'], flags=re.I): gio_th += a['gio']
        dg = [(x['hinh_thuc'] or x['nhom'], x['trong_so'], x['cong_cu']) for x in out['danh_gia'] if k in x['clo']]
        plo = []
        if out['ma_tran']:
            for r in out['ma_tran']['dong']:
                if re.match(r'^CLO\s*0*%d$' % k, r['nhan'].replace(' ', ''), flags=re.I):
                    plo = [f"{p}={v}" for p, v in r['gia_tri'].items() if v]
        doi_chieu.append({'clo': k, 'muc': c['muc'], 'plo': plo, 'buoi': buoi51, 'gio_theo_hoat_dong': acts, 'gio_thuc_hanh_thao_luan': gio_th, 'danh_gia': dg})
    out['doi_chieu'] = doi_chieu
    # ghi chú cơ học (không phải kết luận)
    for x in doi_chieu:
        if not x['buoi']: out['ghi_chu'].append(f"CLO{x['clo']}: không buổi nào trong 5.1 ghi đóng góp.")
        if not x['danh_gia']: out['ghi_chu'].append(f"CLO{x['clo']}: không hình thức đánh giá nào ghi.")
        if x['muc'].startswith(('K3', 'K4', 'K5', 'K6', 'S')) and x['gio_thuc_hanh_thao_luan'] == 0 and x['buoi']:
            out['ghi_chu'].append(f"CLO{x['clo']} ({x['muc']}): các buổi hướng tới CLO này có 0 giờ thực hành/thảo luận; xem 'Tiểu luận, bài tập lớn, thực tế' có bù không.")
        if x['muc'].startswith('MƠ HỒ') or x['muc'].startswith('?'): out['ghi_chu'].append(f"CLO{x['clo']}: động từ mở đầu {x['muc']}.")
        chi_cc = x['danh_gia'] and all(re.search(r'chuyên cần|attendance', h, flags=re.I) for h, _, _ in x['danh_gia'])
        if chi_cc: out['ghi_chu'].append(f"CLO{x['clo']}: chỉ được đánh giá bằng chuyên cần.")
        if len(x['plo']) > 1: out['ghi_chu'].append(f"CLO{x['clo']}: đáp ứng {len(x['plo'])} PLO ({', '.join(x['plo'])}); Hướng dẫn CĐR: mỗi CLO một PLO.")
    return out

def to_md(o):
    L = [f"# Trích xuất đề cương: {o['file']}", '', '## Đầu đề']
    L += [f"- {k}: {v}" for k, v in o['dau_de'].items()]
    L += ['', '## Mô tả học phần', o['mo_ta'] or '(không thấy)', '', '## CLO (mức đoán theo động từ mở đầu)']
    L += [f"- CLO{c['n']} [{c['muc']}]: {c['text']}" for c in o['clo']]
    L += ['', '## Ma trận 3.2']
    if o['ma_tran']:
        L.append('| | ' + ' | '.join(o['ma_tran']['plo']) + ' |'); L.append('|' + '---|' * (len(o['ma_tran']['plo']) + 1))
        for r in o['ma_tran']['dong']: L.append(f"| {r['nhan']} | " + ' | '.join(r['gia_tri'].get(p, '') or '·' for p in o['ma_tran']['plo']) + ' |')
    else: L.append('(không thấy bảng ma trận)')
    L += ['', '## Bảng 5.1', '| Buổi | Nội dung | Hình thức | LT/TH/TL/TH | CLO |', '|---|---|---|---|---|']
    L += [f"| {r['buoi']} | {r['noi_dung'][:90]} | {r['hinh_thuc']} | {'/'.join(str(int(x)) if x == int(x) else str(x) for x in r['gio'])} | {','.join(map(str, r['clo']))} |" for r in o['b51']]
    L += ['', '## Bảng 5.2 (hoạt động có giờ hoặc có nội dung)']
    for b, v in o['b52'].items():
        L.append(f"- Buổi {b} (CLO {','.join(map(str, v['clo']))}):")
        for a in v['hoat_dong']:
            if a['gio'] > 0 or a['noi_dung']: L.append(f"    - {a['loai']} [{a['gio']:g} giờ]: {a['noi_dung'][:160].replace(chr(10), ' / ')}")
    L += ['', '## Bảng đánh giá', '| Nhóm | Hình thức | Nội dung | Công cụ, tiêu chí | CLO | Trọng số |', '|---|---|---|---|---|---|']
    L += [f"| {x['nhom']} | {x['hinh_thuc']} | {x['noi_dung'][:80]} | {x['cong_cu'][:80]} | {','.join(map(str, x['clo']))} | {x['trong_so']} |" for x in o['danh_gia']]
    L += ['', '## Học liệu']
    for k, v in o['hoc_lieu'].items():
        L.append(f"- **{k}**"); L += [f"    - {x}" for x in v]
    L += ['', '## Đối chiếu CLO → PLO → hoạt động dạy học → đánh giá', '| CLO | Mức | PLO (3.2) | Buổi (5.1) | Giờ theo hoạt động (5.2) | Đánh giá (hình thức, trọng số) |', '|---|---|---|---|---|---|']
    for x in o['doi_chieu']:
        acts = '; '.join(f"{k}: {v:g}" for k, v in x['gio_theo_hoat_dong'].items())
        dg = '; '.join(f"{h} {w}" for h, w, _ in x['danh_gia'])
        L.append(f"| CLO{x['clo']} | {x['muc']} | {', '.join(x['plo']) or '·'} | {', '.join(x['buoi']) or '·'} | {acts or '·'} | {dg or '·'} |")
    L += ['', '## Ghi chú cơ học (cần người đọc kết luận)'] + ([f"- {g}" for g in o['ghi_chu']] or ['- (không có)'])
    return '\n'.join(L)

if __name__ == '__main__':
    if len(sys.argv) < 2: print(__doc__); sys.exit(1)
    o = extract(sys.argv[1])
    if '--json' in sys.argv: print(json.dumps(o, ensure_ascii=False, indent=1))
    else: print(to_md(o))
